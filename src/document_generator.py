from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

from src.auto_apply_rules import auto_apply_block_reason, is_auto_apply_blocked
from src.cv_selector import select_cv_profile
from src.database import Database
from src.fit_scorer import CORE_SKILLS, score_job
from src.salary_estimator import estimate_salary_expectation

PLACEHOLDER_MARKERS = [
    "[Your Name]",
    "[Email]",
    "[Phone]",
    "email@example.com",
    "linkedin.com/in/your-profile",
    "Your Name",
    "Replace this template",
    "Profile Summary Template",
]

COVER_LETTER_FORBIDDEN_PHRASES = [
    "cv profile",
    "fit score",
    "fit scoring",
    "best matched",
    "aligns well with",
    "the main value here seems to be",
    "that is the kind of role i am aiming for",
    "ai job apply assistant",
    "application workflow",
    "scoring",
    "automation pipeline",
    "internal logic",
    "document basis",
    "email snippet fallback",
    "role classification",
    "ai reasoning",
    "role family",
]

SAFE_COMPANY_CONTEXTS = [
    "a company using data to improve operational decision-making",
    "a team focused on data, reporting, and business improvement",
    "an organisation where analytics supports better infrastructure and service outcomes",
]

SAFE_COMPANY_MOTIVATIONS = [
    "the opportunity to contribute to practical reporting and analytics work that supports better business decisions",
    "the chance to work in a data-focused environment where reporting, insight, and stakeholder communication are important",
    "the opportunity to apply my background in Power BI, SQL, and analytics to real operational challenges",
]

SAFE_RELEVANT_SKILLS = [
    "Power BI, SQL, reporting automation, data modelling, and stakeholder communication",
    "Python, SQL, Power BI, data workflows, and analytics",
    "business intelligence, dashboard development, SQL reporting, and data analysis",
]

SAFE_GROWTH_AREAS = [
    "modern data platforms and analytics engineering",
    "data engineering, automation, and cloud-based reporting tools",
    "business intelligence, data modelling, and emerging AI-enabled analytics tools",
]

RECRUITER_FACING_CV_BASENAME = "Example_Candidate_CV"


class DocumentGenerator:
    def __init__(self, database: Database, app_root: Path) -> None:
        self.database = database
        self.app_root = Path(app_root)
        self.data_dir = self.app_root / "data"
        self.applications_dir = self.app_root / "applications"
        self.temp_uploads_dir = self.data_dir / "temp_uploads"
        self.cv_config_path = self.data_dir / "cv_profiles.json"
        self.cover_letter_template_path = self.data_dir / "templates" / "master_cover_letter.txt"

    def generate_for_job(self, job_id: int) -> dict[str, Any]:
        job = self.database.get_job_details(job_id)
        if job is None:
            raise ValueError(f"Job {job_id} not found")
        job = dict(job)
        job["document_description"], document_basis = self._document_description(job)
        self.database.set_job_document_basis(job_id, document_basis)

        personal_details = self._load_personal_details()
        profile_text = self._load_profile()
        base_cv_text = self._load_base_cv_text(personal_details)
        application_dir = self.applications_dir / str(job_id)
        application_dir.mkdir(parents=True, exist_ok=True)

        rescored = score_job(
            title=str(job.get("enriched_title") or job.get("title") or ""),
            location=str(job.get("enriched_location") or job.get("location") or ""),
            salary=str(job.get("enriched_salary") or job.get("salary") or ""),
            description=str(job["document_description"] or ""),
        )
        selection = select_cv_profile(
            config_path=self.cv_config_path,
            title=str(job.get("enriched_title") or job.get("title") or ""),
            description=str(job["document_description"] or ""),
            role_family=str(rescored.enrichment.get("role_family") or ""),
            matched_skills=list(rescored.enrichment.get("matched_skills") or []),
        )
        salary_estimate = estimate_salary_expectation(job)

        selected_cv_source = (self.app_root / selection.file_path).resolve() if selection.file_path else Path()
        selected_cv_output = application_dir / "selected_cv.docx"
        tailored_cv_output = application_dir / "tailored_cv.docx"
        tailored_cv_preview = application_dir / "tailored_cv_preview.md"
        cover_letter_docx = application_dir / "cover_letter.docx"
        cover_letter_preview = application_dir / "cover_letter_preview.md"
        cover_letter_md = application_dir / "cover_letter.md"
        interview_prep_path = application_dir / "interview_prep.md"
        demo_recommendation_path = application_dir / "demo_recommendation.md"
        checklist_path = application_dir / "application_checklist.md"
        rationale_path = application_dir / "application_rationale.md"
        recruiter_message_path = application_dir / "recruiter_message.txt"
        selection_criteria_path = application_dir / "selection_criteria.md"
        application_packet_path = application_dir / "application_packet.md"

        if selection.file_path and selected_cv_source.exists() and selected_cv_source.is_file():
            shutil.copy2(selected_cv_source, selected_cv_output)
            shutil.copy2(selected_cv_source, tailored_cv_output)
        else:
            self._write_simple_docx(
                selected_cv_output,
                title=selection.profile_name,
                sections=[("CV profile", f"Using fallback CV profile summary because the configured file was unavailable: {selection.file_path or 'not configured'}")],
            )
            shutil.copy2(selected_cv_output, tailored_cv_output)
        upload_cv_temp_path = self._prepare_recruiter_facing_upload_cv(
            source_cv_path=selected_cv_output,
            selection=selection,
            job_id=job_id,
        )

        cover_letter_text, cover_letter_generation_mode, sanitized, banned_phrase_detected = self._build_cover_letter(job, personal_details, profile_text)
        print(f"[cover_letter] generation_mode={cover_letter_generation_mode}", flush=True)
        print(f"[cover_letter] template_used={self.cover_letter_template_path}", flush=True)
        print("[cover_letter] regenerated=True", flush=True)
        print(f"[cover_letter] banned_phrase_detected={banned_phrase_detected or 'none'}", flush=True)
        print(f"[cv_upload] selected_cv_variant={selection.profile_name}", flush=True)
        print(f"[cv_upload] upload_cv_unique_filename={upload_cv_temp_path.name}", flush=True)
        print(f"[cv_upload] upload_cv_unique_path={upload_cv_temp_path}", flush=True)
        print(f"[cv_upload] upload_cv_temp_path={upload_cv_temp_path}", flush=True)
        print(f"[cv_upload] recruiter_facing_cv_filename={upload_cv_temp_path.name}", flush=True)
        deferred_run_logs: list[tuple[str, str, str, str]] = [
            (
                "document_generation",
                "cover_letter_generated",
                "success",
                f"job_id={job_id} generation_mode={cover_letter_generation_mode} template_used={self.cover_letter_template_path.name} regenerated=True banned_phrase_detected={banned_phrase_detected or 'none'}",
            ),
            (
                "document_generation",
                "prepare_upload_cv",
                "success",
                f"job_id={job_id} selected_cv_variant={selection.profile_name} upload_cv_unique_filename={upload_cv_temp_path.name} upload_cv_unique_path={upload_cv_temp_path} upload_cv_temp_path={upload_cv_temp_path} recruiter_facing_cv_filename={upload_cv_temp_path.name}",
            ),
        ]
        interview_prep_text = self._build_interview_prep(job, rescored.enrichment)
        demo_text = self._build_demo_recommendation(job, rescored.enrichment)
        checklist_text = self._build_checklist(job)
        rationale_text = self._build_rationale(job, rescored.enrichment)
        recruiter_message = self._build_recruiter_message(job, personal_details)
        selection_criteria = self._build_selection_criteria(job, rescored.enrichment)
        cv_preview_text = self._build_cv_preview(selection, rescored.enrichment, personal_details)
        packet_text = self._build_application_packet(
            job=job,
            selection=selection,
            rescored=rescored.enrichment,
            cover_letter_text=cover_letter_text,
            interview_prep_text=interview_prep_text,
            demo_text=demo_text,
            checklist_text=checklist_text,
            apply_url=str(job.get("url") or ""),
        )

        tailored_cv_preview.write_text(cv_preview_text, encoding="utf-8")
        cover_letter_preview.write_text(cover_letter_text, encoding="utf-8")
        cover_letter_md.write_text(cover_letter_text, encoding="utf-8")
        interview_prep_path.write_text(interview_prep_text, encoding="utf-8")
        demo_recommendation_path.write_text(demo_text, encoding="utf-8")
        checklist_path.write_text(checklist_text, encoding="utf-8")
        rationale_path.write_text(rationale_text, encoding="utf-8")
        recruiter_message_path.write_text(recruiter_message, encoding="utf-8")
        selection_criteria_path.write_text(selection_criteria, encoding="utf-8")
        application_packet_path.write_text(packet_text, encoding="utf-8")
        self._write_simple_docx(
            cover_letter_docx,
            title=f"Cover Letter - {job.get('title') or 'Job'}",
            sections=[("Cover letter", cover_letter_text)],
        )

        self._validate_written_documents(
            cv_docx_path=tailored_cv_output,
            text_paths=[
                tailored_cv_preview,
                cover_letter_preview,
                cover_letter_md,
                interview_prep_path,
                demo_recommendation_path,
                checklist_path,
                rationale_path,
                recruiter_message_path,
                selection_criteria_path,
                application_packet_path,
            ],
        )

        self.database.insert_generated_document(job_id, "selected_cv", selected_cv_output, selection.reason)
        self.database.insert_generated_document(job_id, "tailored_cv", tailored_cv_output, selection.reason)
        self.database.insert_generated_document(job_id, "tailored_cv_preview", tailored_cv_preview, cv_preview_text)
        self.database.insert_generated_document(job_id, "cover_letter", cover_letter_md, cover_letter_text)
        self.database.insert_generated_document(job_id, "cover_letter_docx", cover_letter_docx, cover_letter_text)
        self.database.insert_generated_document(job_id, "interview_prep", interview_prep_path, interview_prep_text)
        self.database.insert_generated_document(job_id, "demo_recommendation", demo_recommendation_path, demo_text)
        self.database.insert_generated_document(job_id, "selection_criteria", selection_criteria_path, selection_criteria)
        self.database.insert_generated_document(job_id, "recruiter_message", recruiter_message_path, recruiter_message)
        self.database.insert_generated_document(job_id, "checklist", checklist_path, checklist_text)
        self.database.insert_generated_document(job_id, "application_rationale", rationale_path, rationale_text)
        self.database.insert_generated_document(job_id, "application_packet", application_packet_path, packet_text)

        self.database.update_job_career_assets(
            job_id,
            selected_cv_profile=selection.profile_name,
            selected_cv_path=str(upload_cv_temp_path),
            cv_selection_reason=selection.reason,
            salary_expectation_text=salary_estimate.salary_expectation_text,
            salary_estimate_low=salary_estimate.salary_low,
            salary_estimate_high=salary_estimate.salary_high,
            salary_estimate_reasoning_json=json.dumps(salary_estimate.reasoning),
            fit_category=str(rescored.enrichment.get("fit_category") or ""),
            matched_skills_json=json.dumps(list(rescored.enrichment.get("matched_skills") or [])),
            missing_skills_json=json.dumps(list(rescored.enrichment.get("missing_skills") or [])),
            seniority_level=str(rescored.enrichment.get("seniority_level") or rescored.enrichment.get("seniority") or ""),
            role_family=str(rescored.enrichment.get("role_family") or ""),
            opportunity_quality=str(rescored.enrichment.get("opportunity_quality") or ""),
            risk_flags_json=json.dumps(list(rescored.enrichment.get("risk_flags") or [])),
            cover_letter_path=str(cover_letter_docx),
            cover_letter_preview=cover_letter_text,
            cover_letter_basis=cover_letter_generation_mode,
            interview_prep_md=interview_prep_text,
            interview_prep_path=str(interview_prep_path),
            demo_recommendation_md=demo_text,
            demo_recommendation_path=str(demo_recommendation_path),
            application_packet_path=str(application_packet_path),
            document_basis=document_basis,
        )

        blocked_auto_apply = is_auto_apply_blocked(job)
        next_status = "ready_to_apply" if self._looks_easy_apply(job) and not blocked_auto_apply else "reviewed"
        next_note = auto_apply_block_reason(job) if blocked_auto_apply else "Application materials prepared for human review."
        self.database.update_application_state(job_id, next_status, next_note)
        deferred_run_logs.append(
            ("document_generation", "generate_package", "success", f"Generated reviewed application packet for job {job_id}")
        )
        self.database.log_runs(deferred_run_logs)
        return {
            "job_id": job_id,
            "output_dir": str(application_dir),
            "ready_to_apply": next_status == "ready_to_apply",
            "rationale_path": str(rationale_path),
            "cv_preview_path": str(tailored_cv_preview),
            "document_basis": document_basis,
            "selected_cv_profile": selection.profile_name,
            "selected_cv_path": str(upload_cv_temp_path),
            "upload_cv_temp_path": str(upload_cv_temp_path),
            "salary_expectation_text": salary_estimate.salary_expectation_text,
            "salary_estimate_low": salary_estimate.salary_low,
            "salary_estimate_high": salary_estimate.salary_high,
            "salary_estimate_reasoning": salary_estimate.reasoning,
            "cover_letter_generation_mode": cover_letter_generation_mode,
            "cover_letter_sanitized": sanitized,
            "cover_letter_banned_phrase_detected": banned_phrase_detected or "",
        }

    def regenerate_missing_or_placeholder_documents(self) -> dict[str, int]:
        jobs = self.database.search_jobs(status="All", min_fit_score=0)
        regenerated = 0
        clean = 0
        for _, row in jobs.iterrows():
            job_id = int(row["job_id"])
            if self.job_needs_regeneration(job_id):
                self.generate_for_job(job_id)
                regenerated += 1
            else:
                clean += 1
        return {"regenerated": regenerated, "already_clean": clean}

    def job_needs_regeneration(self, job_id: int) -> bool:
        application_dir = self.applications_dir / str(job_id)
        required = [
            application_dir / "selected_cv.docx",
            application_dir / "tailored_cv_preview.md",
            application_dir / "cover_letter.md",
            application_dir / "interview_prep.md",
            application_dir / "demo_recommendation.md",
            application_dir / "application_packet.md",
        ]
        for path in required:
            if not path.exists():
                return True
        for path in [item for item in required if item.suffix.lower() in {".md", ".txt"}]:
            content = path.read_text(encoding="utf-8", errors="ignore")
            if self._contains_placeholder(content):
                return True
        details = self.database.get_job_details(job_id) or {}
        if self.cover_letter_contains_banned_phrases(str(details.get("cover_letter_preview") or "")):
            return True
        return False

    def cover_letter_contains_banned_phrases(self, text: str) -> str | None:
        lowered = str(text or "").lower()
        for phrase in COVER_LETTER_FORBIDDEN_PHRASES:
            if phrase in lowered:
                return phrase
        return None

    def ensure_current_cover_letter(self, job_id: int, *, force_regenerate: bool = False) -> dict[str, Any]:
        details = self.database.get_job_details(job_id)
        if details is None:
            raise ValueError(f"Job {job_id} not found")
        application_dir = self.applications_dir / str(job_id)
        cover_letter_path = application_dir / "cover_letter.md"
        stored_preview = str(details.get("cover_letter_preview") or "")
        stored_banned = self.cover_letter_contains_banned_phrases(stored_preview)
        file_banned = None
        if cover_letter_path.exists():
            file_banned = self.cover_letter_contains_banned_phrases(cover_letter_path.read_text(encoding="utf-8", errors="ignore"))
        needs_regeneration = (
            force_regenerate
            or not stored_preview.strip()
            or not cover_letter_path.exists()
            or stored_banned is not None
            or file_banned is not None
        )
        if needs_regeneration:
            detected = stored_banned or file_banned or ""
            print(f"[cover_letter] generation_mode=regeneration_requested", flush=True)
            print(f"[cover_letter] template_used={self.cover_letter_template_path}", flush=True)
            print("[cover_letter] regenerated=True", flush=True)
            print(f"[cover_letter] banned_phrase_detected={detected or 'none'}", flush=True)
            self.database.log_run(
                "document_generation",
                "cover_letter_regeneration_requested",
                "warning" if detected else "success",
                f"job_id={job_id} force_regenerate={force_regenerate} banned_phrase_detected={detected or 'none'}",
            )
            self.generate_for_job(job_id)
            details = self.database.get_job_details(job_id)
            if details is None:
                raise ValueError(f"Job {job_id} not found after regeneration")
        return dict(details)

    def invalidate_stale_cover_letters(self) -> dict[str, int]:
        jobs = self.database.search_jobs(status="All", min_fit_score=0)
        scanned = 0
        regenerated = 0
        for _, row in jobs.iterrows():
            job_id = int(row["job_id"])
            scanned += 1
            details = self.database.get_job_details(job_id) or {}
            banned = self.cover_letter_contains_banned_phrases(str(details.get("cover_letter_preview") or ""))
            if banned:
                self.ensure_current_cover_letter(job_id, force_regenerate=True)
                regenerated += 1
        if scanned:
            self.database.log_run(
                "document_generation",
                "invalidate_stale_cover_letters",
                "success",
                f"scanned={scanned} regenerated={regenerated}",
            )
        return {"scanned": scanned, "regenerated": regenerated}

    def _load_personal_details(self) -> dict[str, str]:
        defaults = {
            "name": "Example Candidate",
            "email": "candidate@example.com",
            "phone": "",
            "linkedin": "",
            "location": "Auckland, New Zealand",
        }
        details_path = self.data_dir / "personal_details.json"
        if details_path.exists():
            loaded = json.loads(details_path.read_text(encoding="utf-8"))
            for key in defaults:
                value = str(loaded.get(key, defaults[key])).strip()
                if value:
                    defaults[key] = value
        return defaults

    def _load_profile(self) -> str:
        profile_path = self.data_dir / "profile.md"
        return profile_path.read_text(encoding="utf-8") if profile_path.exists() else ""

    def _load_base_cv_text(self, personal_details: dict[str, str]) -> str:
        text_path = self.data_dir / "base_cv.txt"
        if text_path.exists():
            content = text_path.read_text(encoding="utf-8")
            return self._inject_personal_details(self._strip_placeholders(content), personal_details)
        return self._inject_personal_details("", personal_details)

    def _build_cover_letter(self, job: dict[str, Any], personal_details: dict[str, str], profile_text: str) -> tuple[str, str, bool, str | None]:
        template = self._load_master_cover_letter_template()
        context = self.generate_cover_letter_context(job, profile_text)
        cover_letter = self._render_master_cover_letter(template, context, personal_details)
        sanitized_cover_letter, banned_phrase = self.sanitize_cover_letter(cover_letter)
        was_sanitized = banned_phrase is not None
        if was_sanitized:
            fallback_context = self._fallback_cover_letter_context(job)
            sanitized_cover_letter = self._render_master_cover_letter(template, fallback_context, personal_details)
            sanitized_cover_letter, fallback_banned_phrase = self.sanitize_cover_letter(sanitized_cover_letter)
            if fallback_banned_phrase is not None:
                raise ValueError("Pure template fallback still produced an unsafe cover letter")
            self.database.log_run(
                "document_generation",
                "cover_letter_sanitized",
                "warning",
                f"job_id={job.get('job_id') or job.get('id') or 'unknown'} cover_letter_sanitized=True banned_phrase_detected={banned_phrase or 'none'}",
            )
            return sanitized_cover_letter, "template_fallback_sanitized", True, banned_phrase
        return sanitized_cover_letter, "template_contextualized", False, None

    def generate_cover_letter_context(self, job: dict[str, Any], profile_text: str) -> dict[str, str]:
        job_title = str(job.get("enriched_title") or job.get("title") or "Data Analyst").strip() or "Data Analyst"
        company = str(job.get("enriched_company") or job.get("company") or "the company").strip() or "the company"
        location = str(job.get("enriched_location") or job.get("location") or "Auckland, New Zealand").strip()
        description = str(job.get("document_description") or "").strip()
        combined_text = " ".join(
            [
                job_title,
                company,
                location,
                description,
                str(job.get("source_page_text") or ""),
                str(profile_text or ""),
            ]
        ).lower()

        if any(term in combined_text for term in ["consulting", "client", "advisory", "transformation"]):
            company_context = SAFE_COMPANY_CONTEXTS[1]
            company_motivation = SAFE_COMPANY_MOTIVATIONS[1]
        elif any(term in combined_text for term in ["infrastructure", "utilities", "construction", "engineering", "maintenance", "facilities", "asset"]):
            company_context = SAFE_COMPANY_CONTEXTS[2]
            company_motivation = SAFE_COMPANY_MOTIVATIONS[0]
        elif any(term in combined_text for term in ["platform", "warehouse", "fabric", "azure", "databricks", "snowflake", "pipeline", "etl", "elt"]):
            company_context = SAFE_COMPANY_CONTEXTS[1]
            company_motivation = SAFE_COMPANY_MOTIVATIONS[2]
        else:
            company_context = SAFE_COMPANY_CONTEXTS[0]
            company_motivation = SAFE_COMPANY_MOTIVATIONS[0]

        if any(term in combined_text for term in ["power bi", "stakeholder", "dashboard", "reporting"]):
            relevant_skills = SAFE_RELEVANT_SKILLS[0]
        elif any(term in combined_text for term in ["python", "sql", "data workflow", "etl", "elt"]):
            relevant_skills = SAFE_RELEVANT_SKILLS[1]
        else:
            relevant_skills = SAFE_RELEVANT_SKILLS[2]

        if any(term in combined_text for term in ["fabric", "azure", "databricks", "spark", "pipeline", "etl", "elt", "data engineer", "analytics engineer"]):
            growth_area = SAFE_GROWTH_AREAS[0]
        elif any(term in combined_text for term in ["cloud", "automation", "reporting", "platform"]):
            growth_area = SAFE_GROWTH_AREAS[1]
        else:
            growth_area = SAFE_GROWTH_AREAS[2]

        return {
            "date": self._cover_letter_date(job),
            "company": company,
            "location": location or "Auckland, New Zealand",
            "job_title": job_title,
            "company_context": company_context,
            "company_motivation": company_motivation,
            "relevant_skills": relevant_skills,
            "growth_area": growth_area,
        }

    def _fallback_cover_letter_context(self, job: dict[str, Any]) -> dict[str, str]:
        return {
            "date": self._cover_letter_date(job),
            "company": str(job.get("enriched_company") or job.get("company") or "the company").strip() or "the company",
            "location": str(job.get("enriched_location") or job.get("location") or "Auckland, New Zealand").strip() or "Auckland, New Zealand",
            "job_title": str(job.get("enriched_title") or job.get("title") or "Data Analyst").strip() or "Data Analyst",
            "company_context": SAFE_COMPANY_CONTEXTS[0],
            "company_motivation": SAFE_COMPANY_MOTIVATIONS[0],
            "relevant_skills": SAFE_RELEVANT_SKILLS[0],
            "growth_area": SAFE_GROWTH_AREAS[0],
        }

    def _build_interview_prep(self, job: dict[str, Any], enrichment: dict[str, Any]) -> str:
        matched = list(enrichment.get("matched_skills") or [])
        missing = list(enrichment.get("missing_skills") or [])
        role_family = str(enrichment.get("role_family") or "unknown").replace("_", " ")
        return (
            "# Interview Prep\n\n"
            "## Likely interview questions\n"
            f"- How have you used {matched[0] if matched else 'SQL or Power BI'} in a business-facing setting?\n"
            "- Tell us about a time you had to turn a vague stakeholder request into something measurable and useful.\n"
            f"- Why does this {role_family} role make sense as your next step?\n\n"
            "## Technical topics to revise\n"
            f"- {', '.join(matched[:5]) or 'SQL, Power BI, Python, reporting design, and data quality'}\n\n"
            "## Business and stakeholder questions\n"
            "- How do you prioritise conflicting reporting requests?\n"
            "- How do you validate data before showing it to stakeholders?\n\n"
            "## Gaps to prepare for\n"
            f"- {', '.join(missing) or 'Be ready to explain where your experience is strongest and where you are still building depth.'}\n\n"
            "## Stories from Example Candidate's experience\n"
            "- Power BI dashboard delivery for operational and stakeholder reporting.\n"
            "- SQL and data validation work that improved reporting accuracy and confidence.\n"
            "- Python-based workflow improvements that reduced manual reporting effort.\n\n"
            "## Suggested answer angles\n"
            "- Emphasise practical delivery, trusted outputs, and communication.\n"
            "- Show how business analysis and technical implementation work together in your approach.\n\n"
            "## 30-minute prep plan\n"
            "- Review the job description and match it to 3 relevant examples.\n"
            "- Revise the main tools named in the ad.\n\n"
            "## 2-hour prep plan\n"
            "- Prepare STAR stories, revise likely technical topics, and review the generated demo recommendation.\n\n"
            "## 1-day prep plan\n"
            "- Build or polish a small proof task, rehearse answers aloud, and review company context from the ad."
        )

    def _build_demo_recommendation(self, job: dict[str, Any], enrichment: dict[str, Any]) -> str:
        role_family = str(enrichment.get("role_family") or "unknown")
        if role_family == "data_engineering":
            project = "Fabric or Python mini pipeline"
            tech_stack = "Python, SQLite, optional Fabric or Databricks notebook"
            steps = "- Ingest sample data\n- Clean and validate it\n- Load into a simple mart\n- Expose a small reporting view"
        elif role_family in {"analytics_engineering", "business_intelligence"}:
            project = "Power BI semantic model and stakeholder dashboard"
            tech_stack = "Power BI, SQL, CSV or SQLite source"
            steps = "- Model a small mart\n- Build a semantic layer\n- Deliver a dashboard with business KPIs\n- Add data checks"
        else:
            project = "Streamlit workflow automation demo"
            tech_stack = "Python, Streamlit, SQLite"
            steps = "- Build a small input pipeline\n- Enrich or score records\n- Show a decision dashboard\n- Explain why it improves workflow quality"
        return (
            "# Demo Recommendation\n\n"
            f"## Recommended demo project\n{project}\n\n"
            "## Why this demo helps\nIt shows practical delivery, structured thinking, and the ability to turn messy requirements into usable outputs.\n\n"
            f"## Estimated time\n4-8 hours\n\n## Tech stack\n{tech_stack}\n\n## Steps\n{steps}\n\n"
            "## How to explain it in interview\nFrame it as a small but realistic example of how you approach data quality, stakeholder usefulness, and delivery discipline."
        )

    def _build_checklist(self, job: dict[str, Any]) -> str:
        return (
            "Application checklist:\n"
            "- Review the job page and confirm the role is still open.\n"
            "- Review the selected CV profile and cover letter.\n"
            "- Review interview prep and demo recommendation.\n"
            "- If using Apply Assist, review all fields before submitting. This app does not submit applications automatically.\n"
            "- Manually submit only after review."
        )

    def _build_rationale(self, job: dict[str, Any], enrichment: dict[str, Any]) -> str:
        return (
            "# Application Rationale\n\n"
            f"## Why this role\nThis role aligns with Example Candidate's {str(enrichment.get('role_family') or 'analytics')} direction and strengths in {', '.join(list(enrichment.get('matched_skills') or [])[:4]) or 'Power BI, SQL, Python, and stakeholder delivery'}.\n\n"
            "## Risks or gaps\n"
            f"- {', '.join(list(enrichment.get('risk_flags') or [])) or 'No major blockers from the current information, but depth should still be discussed honestly.'}\n"
        )

    def _build_recruiter_message(self, job: dict[str, Any], personal_details: dict[str, str]) -> str:
        return (
            f"Hi, I'm {personal_details['name']}. I'm interested in the {job.get('title')} opportunity at {job.get('company')}. "
            "My background is strongest in Power BI, SQL, Python, analytics delivery, and automation workflows. "
            "Happy to share my CV and a short tailored note if the role is still open."
        )

    def _build_selection_criteria(self, job: dict[str, Any], enrichment: dict[str, Any]) -> str:
        matched = ", ".join(list(enrichment.get("matched_skills") or [])[:5]) or "Power BI, SQL, Python, reporting, and stakeholder delivery"
        return (
            "Suggested response points:\n"
            f"- Technical match: {matched}\n"
            "- Delivery match: demonstrate a reporting, dashboard, or workflow example.\n"
            "- Stakeholder match: show how you translated requirements into something useful and trusted."
        )

    def _build_cv_preview(self, selection: Any, enrichment: dict[str, Any], personal_details: dict[str, str]) -> str:
        return (
            f"# {personal_details['name']}\n\n"
            f"Selected CV profile: {selection.profile_name}\n\n"
            f"Selection reason: {selection.reason}\n\n"
            f"Role family: {enrichment.get('role_family')}\n"
            f"Fit category: {enrichment.get('fit_category')}\n"
            f"Matched skills: {', '.join(list(enrichment.get('matched_skills') or []))}\n"
        )

    def _build_application_packet(
        self,
        *,
        job: dict[str, Any],
        selection: Any,
        rescored: dict[str, Any],
        cover_letter_text: str,
        interview_prep_text: str,
        demo_text: str,
        checklist_text: str,
        apply_url: str,
    ) -> str:
        return (
            "# Application Packet\n\n"
            f"- Selected CV profile: {selection.profile_name}\n"
            f"- Fit score: {job.get('fit_score')}\n"
            f"- Fit category: {rescored.get('fit_category')}\n"
            f"- Role family: {rescored.get('role_family')}\n"
            f"- Missing skills: {', '.join(list(rescored.get('missing_skills') or []))}\n"
            f"- Apply URL: {apply_url}\n\n"
            "## Cover Letter\n"
            f"{cover_letter_text}\n\n"
            "## Interview Prep\n"
            f"{interview_prep_text}\n\n"
            "## Demo Recommendation\n"
            f"{demo_text}\n\n"
            "## Checklist\n"
            f"{checklist_text}\n"
        )

    def _prepare_recruiter_facing_upload_cv(self, *, source_cv_path: Path, selection: Any, job_id: int) -> Path:
        self.temp_uploads_dir.mkdir(parents=True, exist_ok=True)
        extension = source_cv_path.suffix.lower() or ".docx"
        timestamp_suffix = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        upload_filename = f"{RECRUITER_FACING_CV_BASENAME}_{timestamp_suffix}{extension}"
        upload_cv_temp_path = self.temp_uploads_dir / upload_filename
        if upload_cv_temp_path.exists():
            upload_cv_temp_path.unlink()
        shutil.copy2(source_cv_path, upload_cv_temp_path)
        if source_cv_path.read_bytes() != upload_cv_temp_path.read_bytes():
            raise ValueError(
                f"Recruiter-facing CV upload copy did not match selected variant for job {job_id}: {selection.profile_name}"
            )
        return upload_cv_temp_path

    def _load_master_cover_letter_template(self) -> str:
        if self.cover_letter_template_path.exists():
            return self.cover_letter_template_path.read_text(encoding="utf-8")
        fallback_path = Path(__file__).resolve().parent.parent / "data" / "templates" / "master_cover_letter.txt"
        if fallback_path.exists():
            return fallback_path.read_text(encoding="utf-8")
        raise ValueError(f"Missing master cover letter template: {self.cover_letter_template_path}")

    def _render_master_cover_letter(self, template: str, context: dict[str, str], personal_details: dict[str, str]) -> str:
        cover_letter = template.format(
            date=context["date"],
            company=context["company"],
            location=context["location"],
            job_title=context["job_title"],
            company_context=context["company_context"],
            company_motivation=context["company_motivation"],
            relevant_skills=context["relevant_skills"],
            growth_area=context["growth_area"],
        )
        company = str(context.get("company") or "").strip()
        if company and company.lower() not in cover_letter.lower():
            cover_letter = cover_letter.replace("Hiring Team\n", f"Hiring Team\n{company}\n", 1)
        cover_letter = self._sanitize_text(cover_letter, personal_details)
        return self._normalize_cover_letter_spacing(cover_letter)

    def _cover_letter_date(self, job: dict[str, Any]) -> str:
        imported_at = str(job.get("imported_at") or "").strip()
        if re.match(r"^\d{4}-\d{2}-\d{2}", imported_at):
            year, month, day = imported_at[:10].split("-")
            months = {
                "01": "January", "02": "February", "03": "March", "04": "April",
                "05": "May", "06": "June", "07": "July", "08": "August",
                "09": "September", "10": "October", "11": "November", "12": "December",
            }
            return f"{int(day)} {months.get(month, month)} {year}"
        return "7 May 2026"

    def _write_simple_docx(self, output_path: Path, *, title: str, sections: list[tuple[str, str]]) -> None:
        document = Document()
        document.add_heading(title, level=1)
        for heading, body in sections:
            document.add_heading(heading, level=2)
            for block in str(body or "").split("\n\n"):
                if block.strip():
                    document.add_paragraph(block.strip())
        document.save(output_path)

    def _looks_easy_apply(self, job: dict[str, Any]) -> bool:
        text = " ".join([job.get("title") or "", job.get("source") or "", job.get("document_description") or "", job.get("url") or ""]).lower()
        if any(marker in text for marker in ["workday", "successfactors", "greenhouse", "taleo", "lever", "brassring", "application form"]):
            return False
        return any(marker in text for marker in ["easy apply", "linkedin", "seek apply", "quick apply"])

    def _matched_and_missing_skills(self, job: dict[str, Any], profile: str, base_cv_text: str) -> tuple[list[str], list[str]]:
        profile_text = f"{profile}\n{base_cv_text}".lower()
        job_text = " ".join(
            [
                str(job.get("title") or ""),
                str(job.get("document_description") or ""),
                str(job.get("requirements") or ""),
                str(job.get("responsibilities") or ""),
            ]
        ).lower()
        matched: list[str] = []
        missing: list[str] = []
        for skill in CORE_SKILLS:
            if skill in job_text:
                if skill in profile_text:
                    matched.append(skill)
                else:
                    missing.append(skill)
        return matched, missing

    def _strip_placeholders(self, text: str) -> str:
        cleaned = text or ""
        for marker in PLACEHOLDER_MARKERS:
            cleaned = cleaned.replace(marker, "")
        return cleaned.replace("  ", " ").strip()

    def _sanitize_text(self, text: str, personal_details: dict[str, str]) -> str:
        cleaned = text or ""
        replacements = {
            "[Your Name]": personal_details["name"],
            "Your Name": personal_details["name"],
            "[Email]": personal_details["email"],
            "email@example.com": personal_details["email"],
            "[Phone]": personal_details["phone"] or "",
            "linkedin.com/in/your-profile": personal_details["linkedin"] or "",
            "Replace this template": "",
            "Profile Summary Template": "",
        }
        for marker, replacement in replacements.items():
            cleaned = cleaned.replace(marker, replacement)
        return cleaned.strip()

    def _normalize_cover_letter_spacing(self, text: str) -> str:
        lines = [line.rstrip() for line in str(text or "").splitlines()]
        normalized: list[str] = []
        previous_blank = False
        for line in lines:
            blank = not line.strip()
            if blank and previous_blank:
                continue
            normalized.append(line)
            previous_blank = blank
        return "\n".join(normalized).strip()

    def _validate_cover_letter_text(self, text: str) -> None:
        if self._contains_placeholder(text):
            raise ValueError("Generated cover letter still contains placeholder content")
        lowered = text.lower()
        for phrase in COVER_LETTER_FORBIDDEN_PHRASES:
            if phrase in lowered:
                raise ValueError(f"Generated cover letter contains forbidden internal phrase: {phrase}")

    def sanitize_cover_letter(self, text: str) -> tuple[str, str | None]:
        cleaned = self._normalize_cover_letter_spacing(self._sanitize_text(text, self._load_personal_details()))
        return cleaned, self.cover_letter_contains_banned_phrases(cleaned)

    def _validate_written_documents(self, cv_docx_path: Path, text_paths: list[Path]) -> None:
        for path in text_paths:
            if not path.exists():
                raise ValueError(f"Required generated document missing: {path.name}")
            content = path.read_text(encoding="utf-8", errors="ignore")
            if self._contains_placeholder(content):
                raise ValueError(f"Generated document still contains placeholder content: {path.name}")
        document = Document(cv_docx_path)
        doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        if self._contains_placeholder(doc_text):
            raise ValueError("Generated DOCX CV still contains placeholder content")

    def _inject_personal_details(self, text: str, personal_details: dict[str, str]) -> str:
        if not text.strip():
            details = [personal_details["name"], personal_details["email"], personal_details["location"]]
            if personal_details["phone"]:
                details.append(personal_details["phone"])
            if personal_details["linkedin"]:
                details.append(personal_details["linkedin"])
            return "\n".join(details)
        return text

    def _contains_placeholder(self, text: str) -> bool:
        lowered = text.lower()
        return any(marker.lower() in lowered for marker in PLACEHOLDER_MARKERS)

    def _document_description(self, job: dict[str, Any]) -> tuple[str, str]:
        about_the_job = str(job.get("about_the_job") or "").strip()
        if about_the_job:
            return about_the_job, "about_the_job"
        full_description = str(job.get("full_description") or "").strip()
        if full_description:
            return full_description, "full_description"
        source_page_text = str(job.get("source_page_text") or "").strip()
        if source_page_text:
            return source_page_text, "source_page_text"
        return str(job.get("description") or "").strip(), "email_snippet_fallback"
